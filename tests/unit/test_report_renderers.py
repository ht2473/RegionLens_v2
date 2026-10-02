"""
Проверки перехода «описание отчёта → файл»: файл непуст, опознаётся по сигнатуре,
кириллица и пропуски не ломают форматы.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pytest
from django.conf import settings
from openpyxl import load_workbook

from apps.exports.constants import ExportFormat
from apps.exports.renderers import render
from apps.exports.renderers.fonts import register_report_fonts
from apps.exports.reports.base import (
    COLUMN_INTEGER,
    COLUMN_NUMBER,
    COLUMN_PERCENT,
    COLUMN_TEXT,
    Column,
    ReportDocument,
    Section,
    Table,
    format_value,
)

pytestmark = pytest.mark.unit

# Сигнатуры форматов: XLSX — архив ZIP, PDF опознаётся собственной меткой.
ZIP_SIGNATURE = b"PK\x03\x04"
PDF_SIGNATURE = b"%PDF"


def build_document(*, columns: int = 4, rows: int = 12) -> ReportDocument:
    """Собрать отчёт для проверки вывода."""
    table_columns = [Column("territory", "Субъект", COLUMN_TEXT, width=40)]
    table_columns += [
        Column(f"y{year}", str(2000 + year), COLUMN_NUMBER, width=12) for year in range(columns - 1)
    ]

    table_rows = []
    for index in range(rows):
        row: dict[str, object] = {"territory": f"Регион № {index + 1}"}
        for year in range(columns - 1):
            # Каждое третье значение — пропуск: отчёт обязан показывать их явно.
            row[f"y{year}"] = None if (index + year) % 3 == 0 else 1234.5678 + index
        table_rows.append(row)

    return ReportDocument(
        title="Численность населения, «постоянное» <в среднем за год>",
        subtitle="Значения по субъектам Российской Федерации",
        meta=[("Источник данных", "Росстат"), ("Лицензия", "Creative Commons BY")],
        sections=[
            Section(
                heading="Значения показателя",
                paragraphs=["Прочерк означает отсутствие наблюдения, а не нулевое значение."],
                tables=[Table(columns=table_columns, rows=table_rows, title="Значения")],
            )
        ],
        footer="RegionLens · Кузьмин Евгений Олегович",
    )


class TestDocumentModel:
    """Промежуточное представление отчёта."""

    def test_row_count_sums_all_tables(self) -> None:
        """Число строк отчёта складывается из всех его таблиц."""
        document = build_document(rows=7)
        document.sections[0].tables.append(
            Table(columns=[Column("a", "A")], rows=[{"a": 1}, {"a": 2}])
        )
        assert document.row_count == 9

    @pytest.mark.parametrize(
        ("value", "kind", "expected"),
        [
            (None, COLUMN_TEXT, "—"),
            ("", COLUMN_NUMBER, "—"),
            (1234, COLUMN_INTEGER, "1 234"),
            (0.4567, COLUMN_PERCENT, "45.7 %"),
        ],
    )
    def test_value_formatting(self, value: object, kind: str, expected: str) -> None:
        """
        Пропуск обозначается прочерком, а не пустой ячейкой.

        Пустая ячейка неотличима от ошибки формирования отчёта.
        """
        assert format_value(value, Column("k", "K", kind)) == expected


class TestRenderers:
    """Вывод отчёта в файлы."""

    @pytest.mark.parametrize(
        ("export_format", "signature"),
        [
            (ExportFormat.XLSX, ZIP_SIGNATURE),
            (ExportFormat.PDF, PDF_SIGNATURE),
        ],
    )
    def test_file_is_produced(self, export_format: str, signature: bytes) -> None:
        """Файл собирается и опознаётся по сигнатуре формата."""
        payload = render(build_document(), export_format)

        assert payload.startswith(signature)
        assert len(payload) > 1000

    def test_wide_table_is_split_in_pdf(self) -> None:
        """
        Широкая таблица в PDF разрезается на части и не теряет строк.

        Двадцать пять лет наблюдений не помещаются в страницу ни при каком размере
        шрифта, поэтому части нумеруются, а первый столбец повторяется в каждой.
        """
        from apps.exports.renderers.pdf import _split_columns

        columns = [Column(f"c{index}", str(index)) for index in range(26)]
        chunks = _split_columns(columns)

        assert len(chunks) > 1
        assert all(chunk[0] is columns[0] for chunk in chunks)
        restored = [column for chunk in chunks for column in chunk[1:]]
        assert restored == columns[1:]

    def test_narrow_table_is_not_split(self) -> None:
        """Таблица, помещающаяся на страницу, не разрезается."""
        from apps.exports.renderers.pdf import _split_columns

        columns = [Column(f"c{index}", str(index)) for index in range(5)]
        assert _split_columns(columns) == [columns]

    def test_unknown_format_is_rejected(self) -> None:
        """Неизвестный формат не молчит, а сообщает об ошибке."""
        with pytest.raises(ValueError, match="формат"):
            render(build_document(), "rtf")


class TestColours:
    """Цвета документов берутся из маркеров оформления, светлой темы."""

    RENDERERS = Path(settings.BASE_DIR) / "apps" / "exports" / "renderers"

    @staticmethod
    def light(name: str) -> str:
        """Светлое значение переменной из tokens.css."""
        text = (Path(settings.BASE_DIR) / "static" / "css" / "tokens.css").read_text("utf-8")
        match = re.search(rf"{name}:\s*light-dark\(\s*#([0-9a-fA-F]{{6}})", text)
        assert match, name
        return match.group(1).upper()

    def test_table_header_repeats_the_site(self) -> None:
        """Шапка таблицы документа — те же маркеры, что у шапки таблиц сайта."""
        from apps.exports.renderers.palette import document_colours

        palette = document_colours()

        assert palette.header == self.light("--surface-sunken")
        assert palette.header_text == self.light("--text-secondary")
        assert palette.header_rule == self.light("--border-strong")

    def test_renderers_hold_no_colour_literals(self) -> None:
        """В модулях отрисовки нет записанных цветов: при смене облика документы не отстанут."""
        for path in self.RENDERERS.glob("*.py"):
            source = path.read_text("utf-8")
            assert not re.search(r"""["']#?[0-9A-Fa-f]{6}["']""", source), path.name
            assert "colors.white" not in source, path.name

    def test_xlsx_header_is_painted_from_tokens(self) -> None:
        """Шапка листа Excel закрашена цветом маркера."""
        payload = render(build_document(), ExportFormat.XLSX)
        sheet = load_workbook(io.BytesIO(payload))["Значения"]

        assert sheet["A1"].fill.fgColor.rgb.endswith(self.light("--surface-sunken"))


class TestFonts:
    """Подбор шрифта для PDF."""

    def test_registered_font_supports_cyrillic(self) -> None:
        """Подобранный шрифт содержит кириллицу; без системного шрифта проверка пропускается."""
        regular, bold = register_report_fonts()

        if regular == "Helvetica":
            pytest.skip("В системе не найден шрифт с поддержкой кириллицы")

        assert regular.startswith("RegionLens")
        assert bold
