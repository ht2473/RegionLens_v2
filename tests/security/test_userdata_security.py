"""
Безопасность своих данных: разметка из названий таблицы экранируется, ячейки выгрузок
не становятся формулами, изменение без CSRF и чтение исходника без доступа отклоняются.
"""

from __future__ import annotations

import csv
import io
import shutil
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from apps.userdata.models import Dataset, DatasetVersion

pytestmark = [pytest.mark.integration, pytest.mark.django_db]

SCRIPT = "<script>alert(1)</script>"
FORMULA = '=HYPERLINK("http://example.com","x")'


@pytest.fixture(autouse=True)
def userdata_dir(settings: Any, tmp_path: Path) -> Iterator[Path]:
    settings.USERDATA_DIR = tmp_path / "userdata"
    yield settings.USERDATA_DIR
    from apps.warehouse.duckdb_client import close_connections

    close_connections()
    shutil.rmtree(settings.USERDATA_DIR, ignore_errors=True)


def _hostile_table() -> bytes:
    """Длинная таблица: показатель с разметкой, единица — формула, разрез — команда."""
    lines = ["region;year;indicator;unit;group;value"]
    names = ("Москва", "Республика Татарстан", "Свердловская область", "Новосибирская область")
    for number, name in enumerate(names):
        group = "+cmd" if number % 2 else "@cmd"
        for year in (2022, 2023):
            lines.append(f"{name};{year};{SCRIPT};{FORMULA};{group};12,5")
    return "\n".join(lines).encode("utf-8")


def _built(client: Client) -> Dataset:
    client.post(
        reverse("userdata:upload"),
        {"action": "upload", "file": SimpleUploadedFile("hostile.csv", _hostile_table())},
    )
    dataset = Dataset.objects.latest("created_at")
    version = dataset.current_version
    assert version is not None
    table = version.report["inspection"]["tables"][0]
    client.post(
        reverse("userdata:file", args=[dataset.public_id]),
        {"action": "choose", "table": table["key"], "encoding": ""},
    )
    client.post(reverse("userdata:table", args=[dataset.public_id]), {"action": "next"})
    client.get(reverse("userdata:series", args=[dataset.public_id]))
    response = client.post(
        reverse("userdata:series", args=[dataset.public_id]),
        {
            "title": SCRIPT,
            "source_title": SCRIPT,
            "source_url": "",
            "description": SCRIPT,
            "indicator-1": SCRIPT,
            "title-1": SCRIPT,
            "unit-1": FORMULA,
            "kind-1": "relative",
            "polarity-1": "neutral",
        },
    )
    assert response.status_code == 302, response.content[:400]
    dataset.refresh_from_db()
    return dataset


class TestUserText:
    def test_markup_is_escaped(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        version = DatasetVersion.objects.get(pk=dataset.current_version_id)
        record = version.series.first()
        assert record is not None
        key = f"u:{dataset.code}:{record.code}"
        pages = [
            client.get(reverse("userdata:dataset", args=[dataset.public_id])),
            client.get(reverse("userdata:index")),
            client.get(reverse("maps:choropleth"), {"series": key}),
            client.get(reverse("compare:index"), {"series": key}),
            client.get(reverse("rankings:index"), {"series": key}),
        ]
        for page in pages:
            assert page.status_code == 200
            assert SCRIPT not in page.text
            assert "&lt;script&gt;alert(1)&lt;/script&gt;" in page.text

    def test_exports_do_not_start_formulas(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        response = client.get(reverse("userdata:download", args=[dataset.public_id, "csv"]))
        rows = list(
            csv.reader(io.StringIO(b"".join(response.streaming_content).decode("utf-8-sig")))
        )
        cells = [cell for row in rows[1:] for cell in row]
        assert f"'{FORMULA}" in cells
        assert "'+cmd" in cells
        assert "'@cmd" in cells
        assert FORMULA not in cells
        book = load_workbook(
            io.BytesIO(
                client.get(reverse("userdata:download", args=[dataset.public_id, "xlsx"])).content
            )
        )
        values = [cell.value for row in book.active.iter_rows() for cell in row]
        assert f"'{FORMULA}" in values
        assert all(not (isinstance(value, str) and value.startswith("=")) for value in values)

    def test_document_xlsx_has_no_formulas(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        version = DatasetVersion.objects.get(pk=dataset.current_version_id)
        record = version.series.first()
        assert record is not None
        response = client.get(
            reverse("exports:document"),
            {"kind": "series", "series": f"u:{dataset.code}:{record.code}", "format": "xlsx"},
        )
        book = load_workbook(io.BytesIO(response.content))
        for sheet in book.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type != "f", (sheet.title, cell.coordinate, cell.value)

    def test_documents_keep_names_as_text(self, client: Client, warehouse: Any) -> None:
        # Разметка в названии сломала бы разбор абзаца PDF, формула — ячейку XLSX и CSV.
        dataset = _built(client)
        version = DatasetVersion.objects.get(pk=dataset.current_version_id)
        record = version.series.first()
        assert record is not None
        key = f"u:{dataset.code}:{record.code}"
        for kind in ("series", "ranking"):
            for export_format in ("xlsx", "pdf"):
                response = client.get(
                    reverse("exports:document"),
                    {"kind": kind, "series": key, "format": export_format},
                )
                assert response.status_code == 200, (kind, export_format, response.content[:300])
                if export_format == "pdf":
                    assert response.content.startswith(b"%PDF")
                    continue
                book = load_workbook(io.BytesIO(response.content))
                for sheet in book.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            assert cell.data_type != "f", (kind, cell.coordinate, cell.value)
        response = client.get(reverse("exports:data-csv"), {"kind": "series", "series": key})
        assert response.status_code == 200
        rows = csv.reader(io.StringIO(response.content.decode("utf-8-sig")))
        cells = [cell for row in rows for cell in row]
        assert f"'{FORMULA}" in cells or all(FORMULA not in cell for cell in cells)
        assert not [cell for cell in cells if cell.startswith(("=", "+", "@"))]


class TestRequests:
    def test_delete_needs_csrf(self, warehouse: Any) -> None:
        client = Client(enforce_csrf_checks=True)
        dataset = _built(Client())
        response = client.post(
            reverse("userdata:delete", args=[dataset.public_id]), {"confirm": "1"}
        )
        assert response.status_code in {403, 404}
        assert Dataset.objects.filter(pk=dataset.pk).exists()

    def test_source_file_is_attachment(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        response = client.get(reverse("userdata:download", args=[dataset.public_id, "source"]))
        assert response["Content-Disposition"].startswith("attachment")
        assert response["X-Content-Type-Options"] == "nosniff"

    def test_unknown_download_kind(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        url = reverse("userdata:download", args=[dataset.public_id, "parquet"])
        assert client.get(url).status_code == 404


class TestStudiesAndLinks:
    """Исследования и ссылки: разметка экранируется, токен не хранится, правка — с CSRF."""

    def test_study_markup_is_escaped(self, member_client: Client, warehouse: Any) -> None:
        from apps.userdata.models import Study

        dataset = _built(member_client)
        version = DatasetVersion.objects.get(pk=dataset.current_version_id)
        record = version.series.first()
        assert record is not None
        member_client.post(reverse("userdata:study-new"), {"title": SCRIPT})
        study = Study.objects.get(title=SCRIPT)
        edit = reverse("userdata:study-edit", args=[study.public_id])
        member_client.post(edit, {"action": "add-text", "text": SCRIPT})
        member_client.post(
            reverse("userdata:study-add"),
            {
                "target": "map",
                "query_string": f"series=u:{dataset.code}:{record.code}",
                "study": str(study.public_id),
            },
        )
        study.refresh_from_db()
        card = next(block for block in study.blocks if block["kind"] == "view")
        member_client.post(edit, {"action": "change", "block": card["id"], "title": SCRIPT})
        pages = [
            member_client.get(reverse("userdata:study", args=[study.public_id])),
            member_client.get(reverse("userdata:study-card", args=[study.public_id, card["id"]])),
            member_client.get(reverse("userdata:index")),
            member_client.get(reverse("maps:choropleth")),
        ]
        for page in pages:
            assert page.status_code == 200
            assert SCRIPT not in page.text

    def test_token_is_not_stored(self, member_client: Client, warehouse: Any) -> None:
        from apps.userdata.models import Share

        dataset = _built(member_client)
        member_client.post(reverse("userdata:share-create", args=[dataset.public_id]), {})
        page = member_client.get(reverse("userdata:dataset", args=[dataset.public_id]))
        token = page.text.split("/s/", 1)[1].split('"', 1)[0]
        share = Share.objects.get()
        stored = " ".join(str(value) for value in Share.objects.values_list().get())
        assert token not in stored
        assert share.token_hash != token

    def test_study_edit_needs_csrf(self, member: Any, warehouse: Any) -> None:
        from apps.userdata.models import Study

        study = Study.objects.create(owner=member, title="Исследование")
        client = Client(enforce_csrf_checks=True)
        client.force_login(member)
        response = client.post(
            reverse("userdata:study-edit", args=[study.public_id]),
            {"action": "delete", "confirm": "1"},
        )
        assert response.status_code == 403
        assert Study.objects.filter(pk=study.pk).exists()

    def test_share_needs_csrf(self, member: Any, warehouse: Any) -> None:
        from apps.userdata.models import Share

        owner = Client()
        owner.force_login(member)
        dataset = _built(owner)
        client = Client(enforce_csrf_checks=True)
        client.force_login(member)
        response = client.post(reverse("userdata:share-create", args=[dataset.public_id]), {})
        assert response.status_code == 403
        assert not Share.objects.exists()
