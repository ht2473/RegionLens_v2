"""
Свои данные после описания таблицы: шаг «Показатели», сборка файла набора, слой рядов,
пять видов холста, страница таблицы, выгрузка, документы, «Сохранённое» и доступ.

Склад — синтетический (численность для пересчёта на жителя — из его ВРП); таблицы —
усечённые настоящие файлы ЕБТ (``tests/fixtures/userdata``).
"""

from __future__ import annotations

import csv
import io
from datetime import timedelta
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

import duckdb
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.userdata import extract, indicators, jobs, scope
from apps.userdata.models import Dataset, DatasetVersion
from apps.warehouse import queries
from apps.warehouse.queries.sources import population_table
from tests.support.userdata import FIXTURES, userdata_dir  # noqa: F401 — приспособление модуля
from tests.support.userdata import build as _build
from tests.support.userdata import built as _built
from tests.support.userdata import describe as _describe
from tests.support.userdata import key_of as _key
from tests.support.userdata import upload as _upload
from tests.support.userdata import version_of as _version

pytestmark = pytest.mark.integration

VIEWS = (
    "maps:choropleth",
    "compare:index",
    "rankings:index",
    "surface:distribution",
    "surface:table",
)


@pytest.mark.django_db
class TestSeriesStep:
    """Шаг «Показатели»: подсказки по единице, перечень показателей, сборка и переход на карту."""

    def test_page_lists_indicators(self, client: Client, warehouse: Any) -> None:
        dataset = _upload(client, "environment.csv")
        _describe(client, dataset)
        page = client.get(reverse("userdata:series", args=[dataset.public_id]))
        assert page.status_code == 200
        assert "Население в городах с высоким" in page.text
        assert 'name="indicator-1"' in page.text
        version = _version(dataset)
        assert extract.is_current(version)
        # 774 значения, из них 73 — код 8888 («в регионе не ведутся наблюдения»).
        assert version.report["extract"]["values"] == 701

    def test_needs_description_first(self, client: Client, warehouse: Any) -> None:
        dataset = _upload(client, "environment.csv")
        response = client.get(reverse("userdata:series", args=[dataset.public_id]))
        assert response.status_code == 302
        assert response["Location"].endswith(f"/own-data/{dataset.public_id}/table/")

    def test_build_opens_study_with_map(self, client: Client, warehouse: Any) -> None:
        from apps.userdata import studies
        from apps.userdata.models import Study

        dataset = _upload(client, "environment.csv")
        _describe(client, dataset)
        response = _build(client, dataset)
        assert response.status_code == 302
        study = Study.objects.get()
        location = urlparse(response["Location"])
        assert location.path == reverse("userdata:study", args=[study.public_id])
        key = parse_qs(location.query)["series"][0]
        assert key.startswith(f"u:{dataset.code}:")
        (block,) = studies.blocks_of(study)
        assert (block["target"], block["parameters"]["series"]) == ("map", key)
        assert study.title == dataset.title
        # Пересборка той же таблицы ведёт в то же исследование, без новой карточки.
        again = _build(client, dataset)
        assert urlparse(again["Location"]).path == location.path
        assert len(studies.blocks_of(Study.objects.get())) == 1
        version = _version(dataset)
        assert version.state == DatasetVersion.State.BUILT
        assert version.regions_count == 77
        assert (version.first_year, version.last_year) == (2014, 2022)
        assert version.data_path is not None and version.data_path.exists()
        dataset.refresh_from_db()
        assert dataset.state == Dataset.State.READY

    def test_pasted_word_table_opens_study(self, client: Client, warehouse: Any) -> None:
        first = (FIXTURES / "regions_population_part1.txt").read_text(encoding="utf-8")
        client.post(reverse("userdata:upload"), {"action": "paste", "text": first})
        dataset = Dataset.objects.latest("created_at")
        second = (FIXTURES / "regions_population_part2.txt").read_text(encoding="utf-8")
        url = reverse("userdata:file", args=[dataset.public_id])
        client.post(url, {"action": "append", "text": second})
        client.post(url, {"action": "choose", "table": "", "encoding": ""})
        _describe(client, dataset)
        response = _build(client, dataset)
        assert response.status_code == 302
        assert "/own-data/studies/" in urlparse(response["Location"]).path
        version = _version(dataset)
        assert version.regions_count == 85
        assert (version.first_year, version.last_year) == (2010, 2023)

    def test_handwritten_table_opens_study(self, client: Client, warehouse: Any) -> None:
        # Таблица «от руки»: сокращённые и разговорные названия, запятая в числах.
        rows = [
            "Регион;2022;2023",
            "Респ. Татарстан;41,2;42,0",
            "Башкирия;38,5;39,1",
            "г.Москва;55,0;56,3",
            "СПб;51,7;52,2",
            "Московская обл;44,1;45,0",
            "ХМАО — Югра;47,3;47,9",
            "ЯНАО;49,8;50,4",
            "Пермский кр;36,2;36,9",
            "Кузбасс;33,4;34,0",
            "Приморье;35,6;36,1",
        ]
        content = "\n".join(rows).encode("utf-8")
        client.post(
            reverse("userdata:upload"),
            {"action": "upload", "file": SimpleUploadedFile("доля.csv", content)},
        )
        dataset = Dataset.objects.latest("created_at")
        client.post(
            reverse("userdata:file", args=[dataset.public_id]),
            {"action": "choose", "table": "", "encoding": ""},
        )
        _describe(client, dataset)
        response = _build(client, dataset)
        assert response.status_code == 302
        assert "/own-data/studies/" in urlparse(response["Location"]).path
        version = _version(dataset)
        assert version.regions_count == 10
        assert (version.first_year, version.last_year) == (2022, 2023)

    def test_description_never_stops_the_build(self, client: Client, warehouse: Any) -> None:
        # Необязательное — не помеха: пустое название остаётся прежним, адрес без https://
        # дополняется, текст вместо адреса уходит в источник.
        from apps.userdata.forms import DatasetMetaForm

        dataset = _upload(client, "environment.csv")
        blank = {"title": "", "source_title": "", "description": ""}
        form = DatasetMetaForm({**blank, "source_url": "Росстат, бюллетень"}, instance=dataset)
        assert form.is_valid(), form.errors
        assert form.cleaned_data["title"] == dataset.title
        assert form.cleaned_data["source_url"] == ""
        assert form.cleaned_data["source_title"] == "Росстат, бюллетень"
        form = DatasetMetaForm(
            {**blank, "source_title": "Росстат", "source_url": "rosstat.gov.ru/folder/210"},
            instance=dataset,
        )
        assert form.is_valid(), form.errors
        assert form.cleaned_data["source_url"] == "https://rosstat.gov.ru/folder/210"
        form = DatasetMetaForm(
            {**blank, "source_title": "Росстат", "source_url": "сборник 2023"}, instance=dataset
        )
        assert form.is_valid(), form.errors
        assert form.cleaned_data["source_title"] == "Росстат, сборник 2023"

    def test_table_returns_to_its_study(self, client: Client, warehouse: Any) -> None:
        # Куда попадёт таблица, сказано до загрузки; исследование помнится для этой таблицы,
        # и следующая, загруженная не из него, уходит в новое.
        from apps.userdata.models import Study

        client.post(reverse("userdata:study-new"), {"title": "К курсовой"})
        study = Study.objects.get()
        page = client.get(reverse("userdata:upload"), {"study": str(study.public_id)})
        assert "добавится картой в исследование «К курсовой»" in page.text
        assert f'name="study" value="{study.public_id}"' in page.text
        content = (FIXTURES / "environment.csv").read_bytes()
        client.post(
            reverse("userdata:upload"),
            {
                "action": "upload",
                "study": str(study.public_id),
                "file": SimpleUploadedFile("environment.csv", content),
            },
        )
        dataset = Dataset.objects.latest("created_at")
        client.post(
            reverse("userdata:file", args=[dataset.public_id]),
            {"action": "choose", "table": "", "encoding": ""},
        )
        _describe(client, dataset)
        step = client.get(reverse("userdata:series", args=[dataset.public_id]))
        assert "добавится в исследование «К курсовой»" in step.text
        response = _build(client, dataset)
        assert urlparse(response["Location"]).path == reverse(
            "userdata:study", args=[study.public_id]
        )
        opened = client.get(response["Location"])
        assert "добавлена в исследование «К курсовой»" in opened.text

        assert "в новом исследовании" in client.get(reverse("userdata:upload")).text
        other = _upload(client, "crime_wide.csv")
        _describe(client, other)
        again = _build(client, other)
        assert urlparse(again["Location"]).path != reverse("userdata:study", args=[study.public_id])
        assert Study.objects.count() == 2

    def test_rows_under_region_headings_become_a_slice(
        self, client: Client, warehouse: Any
    ) -> None:
        # Вёрстка ФТС: регион — строка без чисел, под ним «Всего» и «Микро предприятия».
        rows = ["Федеральные округа и субъекты РФ;2020 год;2021 год", "ИТОГО;;"]
        rows += ["Всего;10;12", "Микро предприятия;3;4"]
        for region in ("Белгородская область", "Брянская область", "Курская область"):
            rows += [f"{region};;", "Всего;5;6", "Микро предприятия;1;2"]
        content = "\n".join(rows).encode("utf-8")
        client.post(
            reverse("userdata:upload"),
            {"action": "upload", "file": SimpleUploadedFile("экспорт.csv", content)},
        )
        dataset = Dataset.objects.latest("created_at")
        client.post(
            reverse("userdata:file", args=[dataset.public_id]),
            {"action": "choose", "table": "", "encoding": ""},
        )
        page = client.get(reverse("userdata:table", args=[dataset.public_id]))
        assert "Разрез строками под регионом" in page.text
        _describe(client, dataset)
        _build(client, dataset)
        version = _version(dataset)
        assert version.regions_count == 3
        parts = {
            tuple(value for _header, value in item["slices"])
            for item in version.report["extract"]["series"]
        }
        assert parts == {("Всего",), ("Микро предприятия",)}

    def test_table_by_districts_is_explained(self, client: Client, warehouse: Any) -> None:
        rows = [
            "Наименование федерального округа;Объём кредитов",
            "Центральный федеральный округ;100825",
            "Северо-Западный федеральный округ;12510",
            "Южный федеральный округ;445",
        ]
        content = "\n".join(rows).encode("utf-8")
        client.post(
            reverse("userdata:upload"),
            {"action": "upload", "file": SimpleUploadedFile("округа.csv", content)},
        )
        dataset = Dataset.objects.latest("created_at")
        client.post(
            reverse("userdata:file", args=[dataset.public_id]),
            {"action": "choose", "table": "", "encoding": ""},
        )
        page = client.get(reverse("userdata:table", args=[dataset.public_id]))
        assert "В таблице федеральные округа, а не субъекты" in page.text

    def test_unit_hints(self) -> None:
        assert indicators.guess_kind("Число умерших", "человек") == indicators.SUM
        assert indicators.guess_kind("Доля", "процент") == indicators.RELATIVE
        assert (
            indicators.guess_kind("Ожидаемая продолжительность жизни", "лет") == indicators.RELATIVE
        )
        assert indicators.guess_kind("Валовой продукт", "млн руб.") == indicators.SUM
        assert indicators.guess_kind("Средняя заработная плата", "рублей") == indicators.RELATIVE
        assert indicators.guess_kind("Заболеваемость", "на 100 000 человек") == indicators.RELATIVE
        # Слова, похожие на признаки средней и доли, — у сумм (найдено первым взглядом).
        for name, unit in (
            ("Зарегистрировано преступлений средней тяжести", "единиц"),
            ("Среднесписочная численность работников малого и среднего бизнеса", ""),
            ("Преступления в отношении иностранных граждан", "единиц"),
            ("Закуплено годовых курсов лечения", "годовых курсов"),
        ):
            assert indicators.guess_kind(name, unit) == indicators.SUM, name
        assert indicators.guess_kind("Отношение доходов к прожиточному минимуму", "") == (
            indicators.RELATIVE
        )
        assert indicators.guess_kind("Ожидаемая продолжительность", "года") == indicators.RELATIVE

    def test_date_in_title_is_not_per_capita(self) -> None:
        # «на 1 января» — дата, а не «на 1 жителя»: численность остаётся суммой.
        for name, unit in (
            ("Численность населения на 1 января", ""),
            ("Численность населения на 1 января 2024 года", "тыс. человек"),
            ("Задолженность по ипотечным кредитам на 01.01.2007", "млн руб."),
            ("Число организаций на 10-е марта", ""),
        ):
            assert indicators.guess_kind(name, unit) == indicators.SUM, name
            assert indicators.kind_is_clear(name, unit), name
        # Признак «на 100 жителей» рядом с датой по-прежнему узнаётся.
        assert (
            indicators.guess_kind("Врачи на 1 января", "на 10 000 человек населения")
            == indicators.RELATIVE
        )
        # Дата без других признаков вид не подсказывает — его спросят.
        assert not indicators.kind_is_clear("Ипотека на 1 января 2007 года", "")

    def test_plain_rubles_are_per_person(self) -> None:
        # Рубли без разряда — сумма на человека (прожиточный минимум, размер доплаты).
        name = "Величина прожиточного минимума пенсионера"
        assert indicators.guess_kind(name, "руб.") == indicators.RELATIVE
        assert indicators.kind_is_clear(name, "руб.")
        assert indicators.guess_kind("Объём кредитов", "млн руб.") == indicators.SUM

    def test_column_header_gives_name_and_unit(self) -> None:
        cases = (
            ("Средний размер доплаты, руб.", ("Средний размер доплаты", "руб.")),
            (
                "Численность лиц, которым установлена доплата, чел.",
                ("Численность лиц, которым установлена доплата", "чел."),
            ),
            ("Выбросы (тыс. тонн)", ("Выбросы", "тыс. тонн")),
            ("Доля, %", ("Доля", "%")),
            ("Число лет до пенсии", ("Число лет до пенсии", "")),
        )
        for header, expected in cases:
            assert extract.header_unit(header) == expected, header
        # Единица из своего столбца сильнее подписи.
        assert extract.header_unit("Расходы, руб.", "млн рублей") == ("Расходы", "млн рублей")
        # Единица в конце общей подписи составной шапки.
        assert extract.header_unit("Объём отгрузки, млрд руб. | добыча") == (
            "Объём отгрузки | добыча",
            "млрд руб.",
        )

    def test_word_header_is_tidied(self) -> None:
        from apps.userdata.recognize import tidy_header

        assert tidy_header("обрабаты- вающие производства") == "обрабатывающие производства"
        assert tidy_header("деятельности организа- ций4), млрд руб.") == (
            "деятельности организаций, млрд руб."
        )
        # Сноски подряд снимаются вместе, пробел перед пояснением в скобках остаётся.
        assert tidy_header("организаций2);3), млн руб.") == "организаций, млн руб."
        assert tidy_header("населения1)(на конец года)") == "населения (на конец года)"
        # Союз после дефиса и промежуток месяцев — не перенос слова.
        assert tidy_header("сельско- и лесохозяйственная") == "сельско- и лесохозяйственная"
        assert tidy_header("январь- май") == "январь- май"
        assert tidy_header("июнь- сентябрь") == "июнь- сентябрь"
        assert tidy_header("Ввод жилья, тыс. м2 общей площади") == (
            "Ввод жилья, тыс. м2 общей площади"
        )

    def test_word_paste_gives_clean_names(self, client: Client, warehouse: Any) -> None:
        # Вставка из Word: шапка в две строки, «в том числе», переносы слов и сноска.
        text = (FIXTURES / "word_paste_two_level.tsv").read_text(encoding="utf-8")
        client.post(reverse("userdata:upload"), {"action": "paste", "text": text})
        dataset = Dataset.objects.latest("created_at")
        client.post(
            reverse("userdata:file", args=[dataset.public_id]),
            {"action": "choose", "table": "", "encoding": ""},
        )
        # Года в таблице нет — его называет человек.
        _describe(client, dataset, {"year": "2023"})
        client.get(reverse("userdata:series", args=[dataset.public_id]))
        found = {item.name: item for item in indicators.indicators_of(_version(dataset))}
        names = " ".join(found)
        assert found
        assert "- " not in names.replace(" - ", " "), list(found)
        assert "в том числе" not in names, list(found)
        assert "4)" not in names, list(found)
        plant = found["Продукция сельского хозяйства | растениеводства"]
        assert plant.unit == "млрд руб."
        mining = next(name for name in found if name.endswith("| обрабатывающие производства"))
        assert "млрд руб." not in mining
        assert found[mining].unit == "млрд руб."

    def test_caption_is_not_part_of_column_names(self, client: Client, warehouse: Any) -> None:
        # Пример СФР: над шапкой — «Приложение 4» и заголовок, показатели — в столбцах.
        path = Path("data/examples/sfr-pension-supplement-2023.csv")
        dataset = _upload(client, "sfr.csv", path.read_bytes())
        _describe(client, dataset)
        client.get(reverse("userdata:series", args=[dataset.public_id]))
        found = {item.name: item for item in indicators.indicators_of(_version(dataset))}
        assert not any(name.startswith("Приложение") for name in found), list(found)
        minimum = found["Величина прожиточного минимума пенсионера в субъекте Российской Федерации"]
        assert minimum.unit == "руб."
        assert minimum.kind == indicators.RELATIVE
        assert found[
            "Численность лиц, которым установлена федеральная социальная доплата"
        ].unit == ("чел.")

    def test_extraction_failure_is_explained(self, client: Client, warehouse: Any) -> None:
        dataset = _upload(client, "environment.csv")
        # Человек снял роль у столбца регионов: таблицу не разобрать, и об этом сказано.
        _describe(client, dataset, {"role-5": "skip"})
        page = client.get(reverse("userdata:series", args=[dataset.public_id]))
        assert "Разобрать таблицу не удалось" in page.text
        assert "Столбец с регионами не найден" in page.text


@pytest.mark.django_db
class TestBuild:
    """Файл набора: схема склада, витрины, пересчёт на жителя, вложенные области."""

    def test_file_has_warehouse_schema_and_marts(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        version = _version(dataset)
        connection = duckdb.connect(str(version.data_path), read_only=True)
        try:
            tables = {row[0] for row in connection.execute("SHOW TABLES").fetchall()}
            assert {"fact_observation", "dim_series", "mart_rank", "mart_series_stats"} <= tables
            assert connection.execute("SELECT count(*) FROM mart_rank").fetchone()[0] > 650
            levels = dict(
                connection.execute(
                    "SELECT territory_level, count(*) FROM fact_observation GROUP BY 1"
                ).fetchall()
            )
        finally:
            connection.close()
        assert levels["region"] >= 765
        assert levels["country"] == 9

    def test_per_capita_uses_warehouse_population(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client, kind=indicators.SUM, per=("per100000", "perkm2"))
        version = _version(dataset)
        records = {record.derived: record for record in version.series.all()}
        assert set(records) == {"", "per100000", "perkm2"}
        assert records["per100000"].kind == indicators.RELATIVE
        population = population_table()
        connection = duckdb.connect(str(version.data_path), read_only=True)
        try:
            base = dict(
                connection.execute(
                    "SELECT territory_code || ':' || year, value FROM fact_observation "
                    "WHERE series_key = ? AND value IS NOT NULL",
                    [_key(dataset, records[""])],
                ).fetchall()
            )
            derived = connection.execute(
                "SELECT territory_code, year, value, flags FROM fact_observation "
                "WHERE series_key = ? AND territory_code = 'RU-TA'",
                [_key(dataset, records["per100000"])],
            ).fetchall()
        finally:
            connection.close()
        assert derived
        last = max(population["RU-TA"])
        for code, year, value, flags in derived:
            denominator = population[code][min(year, last)]
            assert value == pytest.approx(base[f"{code}:{year}"] / denominator * 100_000)
            # Рассчитано; после последнего года численности — знаменатель заморожен.
            assert flags & 2
            assert bool(flags & 4) == (year > last)

    def test_nested_alone_for_sums(self, client: Client, warehouse: Any) -> None:
        dataset = _upload(client, "crime_wide.csv")
        _describe(client, dataset)
        _build(client, dataset, kind=indicators.SUM, per=())
        version = _version(dataset)
        record = version.series.order_by("order").first()
        assert record is not None
        connection = duckdb.connect(str(version.data_path), read_only=True)
        try:
            rows = dict(
                connection.execute(
                    "SELECT territory_code, value FROM fact_observation "
                    "WHERE series_key = ? AND year = 2020 AND territory_code IN "
                    "('RU-ARK-AGG', 'RU-NEN', 'RU-ARK')",
                    [_key(dataset, record)],
                ).fetchall()
            )
            flags = connection.execute(
                "SELECT flags FROM fact_observation WHERE series_key = ? AND year = 2020 "
                "AND territory_code = 'RU-ARK'",
                [_key(dataset, record)],
            ).fetchone()
        finally:
            connection.close()
        assert rows["RU-ARK"] == pytest.approx(rows["RU-ARK-AGG"] - rows["RU-NEN"])
        assert flags is not None and flags[0] & 2

    def test_nested_alone_stays_empty_for_relative(self, client: Client, warehouse: Any) -> None:
        dataset = _upload(client, "crime_wide.csv")
        _describe(client, dataset)
        _build(client, dataset, kind=indicators.RELATIVE)
        version = _version(dataset)
        assert "RU-ARK" in version.report["build"]["alone_missing"]
        page = client.get(reverse("userdata:dataset", args=[dataset.public_id]), {"tab": "checks"})
        assert "Для долей и средних область без округов не вычислить" in page.text
        # Число замечаний — у вкладки словами, а не точкой без подписи.
        assert '<span class="tabs__note">замечаний: ' in page.text

    def test_codes_survive_rebuild(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        version = _version(dataset)
        first = {record.code for record in version.series.all()}
        file_name = version.data_file
        _build(client, dataset)
        version.refresh_from_db()
        assert {record.code for record in version.series.all()} == first
        assert version.data_file != file_name
        assert not (version.directory / file_name).exists()

    def test_large_table_builds_in_background(
        self, client: Client, warehouse: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        launched: list[tuple[int, str]] = []
        monkeypatch.setattr(jobs, "INLINE_TABLE_BYTES", 10)
        monkeypatch.setattr(jobs, "INLINE_OBSERVATIONS", 10)

        def spawn(arguments: list[str]) -> bool:
            launched.append((int(arguments[1]), arguments[2]))
            return True

        monkeypatch.setattr(jobs, "_spawn", spawn)
        dataset = _upload(client, "environment.csv")
        _describe(client, dataset)
        page = client.get(reverse("userdata:series", args=[dataset.public_id]))
        assert "Разбираем таблицу…" in page.text
        version = _version(dataset)
        assert launched == [(version.pk, jobs.EXTRACT)]
        jobs.run(version, jobs.EXTRACT)
        status = client.get(reverse("userdata:series-status", args=[dataset.public_id]))
        assert status.status_code == 204
        response = _build(client, dataset)
        assert response["Location"].endswith(f"/own-data/{dataset.public_id}/series/")
        assert launched[-1] == (version.pk, jobs.BUILD)
        assert "Собираем таблицу…" in client.get(response["Location"]).text
        jobs.run(version, jobs.BUILD)
        status = client.get(
            reverse("userdata:series-status", args=[dataset.public_id]), {"stage": jobs.BUILD}
        )
        assert status.status_code == 204
        assert "/own-data/studies/" in status["HX-Redirect"]

    def test_queue_waits_for_free_slot(
        self, client: Client, warehouse: Any, monkeypatch: pytest.MonkeyPatch, settings: Any
    ) -> None:
        settings.USERDATA_PARALLEL_JOBS = 0
        monkeypatch.setattr(jobs, "INLINE_TABLE_BYTES", 10)
        monkeypatch.setattr(jobs, "_spawn", lambda arguments: True)
        dataset = _upload(client, "environment.csv")
        _describe(client, dataset)
        page = client.get(reverse("userdata:series", args=[dataset.public_id]))
        assert "Ждём очереди…" in page.text
        assert jobs.stage_state(_version(dataset), jobs.EXTRACT) == jobs.PENDING


@pytest.mark.django_db
class TestCanvas:
    """Пять видов холста, ссылка на вид и «Сохранить вид» — на рядах набора."""

    def test_five_views(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        key = _key(dataset, record)
        for name in VIEWS:
            page = client.get(reverse(name), {"series": key})
            assert page.status_code == 200, name
            assert "видно только вам" in page.text
            assert "Население в городах с высоким" in page.text
            assert f"/own-data/{dataset.public_id}/?show={record.code}" in page.text
            # У ряда своей таблицы нет запроса к программному интерфейсу.
            assert "/api/v1/observations/?series=u" not in page.text

    def test_picker_keeps_only_open_own_series(self, client: Client, warehouse: Any) -> None:
        # Свои ряды живут в лаборатории: на холсте в выборе — только открытый свой ряд.
        dataset = _built(client)
        assert f"u:{dataset.code}:" not in client.get(reverse("maps:choropleth")).text
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        page = client.get(reverse("maps:choropleth"), {"series": _key(dataset, record)})
        assert "data-own" in page.text
        assert f'<option value="{_key(dataset, record)}"' in page.text

    def test_default_year_is_last_full(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        page = client.get(reverse("rankings:index"), {"series": _key(dataset, record)})
        assert page.context["year"] == 2022

    def test_save_view(self, member_client: Client, member: Any, warehouse: Any) -> None:
        from apps.workspace.models import SavedQuery

        dataset = _built(member_client)
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        key = _key(dataset, record)
        member_client.post(
            reverse("workspace:query-save"),
            {"title": "Мой вид", "target": "map", "query_string": f"series={key}&year=2020"},
        )
        query = SavedQuery.objects.get(user=member)
        assert query.series_keys == [key]
        assert "Население в городах" in member_client.get(reverse("workspace:saved")).text
        opened = member_client.get(query.url)
        assert opened.status_code == 200
        # Таблица удалена — вид помечен исчезнувшим, а сам вид отвечает 404.
        dataset.refresh_from_db()
        member_client.post(reverse("userdata:delete", args=[dataset.public_id]), {"confirm": "1"})
        saved = member_client.get(reverse("workspace:saved"))
        assert "ряд удалённой таблицы" in saved.text
        assert member_client.get(query.url).status_code == 404


@pytest.mark.django_db
class TestAccess:
    """Чужая таблица — 404 везде: холст, страница, выгрузка, документ; интерфейс ключи не берёт."""

    def test_other_session(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        key = _key(dataset, record)
        stranger = Client()
        for name in VIEWS:
            assert stranger.get(reverse(name), {"series": key}).status_code == 404, name
        assert (
            stranger.get(reverse("userdata:dataset", args=[dataset.public_id])).status_code == 404
        )
        for kind in ("csv", "xlsx", "source"):
            url = reverse("userdata:download", args=[dataset.public_id, kind])
            assert stranger.get(url).status_code == 404
        document = stranger.get(
            reverse("exports:document"), {"kind": "series", "series": key, "format": "csv"}
        )
        assert document.status_code == 404
        delete = stranger.post(
            reverse("userdata:delete", args=[dataset.public_id]), {"confirm": "1"}
        )
        assert delete.status_code == 404
        assert Dataset.objects.filter(pk=dataset.pk).exists()

    def test_other_account(self, member_client: Client, make_user: Any, warehouse: Any) -> None:
        dataset = _built(member_client)
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        other = Client()
        other.force_login(make_user(email="other@example.com"))
        assert (
            other.get(reverse("maps:choropleth"), {"series": _key(dataset, record)}).status_code
            == 404
        )

    def test_layer_without_scope_sees_nothing(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        key = _key(dataset, record)
        assert queries.year_counts(key) == {}
        assert queries.series_values([key], ["RU-TA"]) == {}
        with scope.reading_as(fingerprint=dataset.guest_key):
            assert sorted(queries.year_counts(key)) == list(range(2014, 2023))

    def test_expired_guest_table(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        Dataset.objects.filter(pk=dataset.pk).update(expires_at=timezone.now() - timedelta(hours=1))
        assert client.get(reverse("userdata:dataset", args=[dataset.public_id])).status_code == 404

    def test_api_rejects_user_keys(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        key = _key(dataset, record)
        for name in ("api:v1:observations", "api:v1:rankings"):
            response = client.get(reverse(name), {"series": key})
            assert response.status_code == 400, name


@pytest.mark.django_db
class TestDatasetPage:
    """Страница таблицы, раздел, выгрузка, документы и удаление."""

    def test_page(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client, kind=indicators.SUM, per=("per100000",))
        page = client.get(reverse("userdata:dataset", args=[dataset.public_id]))
        assert page.status_code == 200
        text = page.text
        assert '<meta name="robots" content="noindex, nofollow">' in text
        assert "на 100 000 жителей" in text
        # Пересчёт на жителей есть — о суммах не предупреждают; территории — на «Проверках».
        checks = client.get(
            reverse("userdata:dataset", args=[dataset.public_id]), {"tab": "checks"}
        )
        assert "Абсолютные величины" not in checks.text
        assert "Субъектов со значениями: 77 из 85." in checks.text
        bare = _built(client, kind=indicators.SUM, per=())
        warned = client.get(reverse("userdata:dataset", args=[bare.public_id]), {"tab": "checks"})
        assert "Абсолютные величины" in warned.text
        section = client.get(reverse("userdata:index"))
        assert dataset.title in section.text

    def test_download_long_table(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        response = client.get(reverse("userdata:download", args=[dataset.public_id, "csv"]))
        assert response.status_code == 200
        assert response["Content-Disposition"].startswith("attachment")
        text = b"".join(response.streaming_content).decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        assert rows[0][:4] == [
            "Территория",
            "Код территории (ISO 3166-2 у субъектов)",
            "ОКАТО",
            "Уровень",
        ]
        # Строка на наблюдение: и со значением, и с пропуском (код 8888 — «нет данных»).
        version = _version(dataset)
        assert len(rows) - 1 == version.report["extract"]["observations"]
        assert sum(1 for row in rows[1:] if "нет данных" in row) == 73
        moscow = next(row for row in rows if row[1] == "RU-MOW" and row[4] == "2020")
        assert moscow[0] == "Москва"
        workbook = client.get(reverse("userdata:download", args=[dataset.public_id, "xlsx"]))
        sheet = load_workbook(io.BytesIO(workbook.content)).active
        assert sheet.max_row == len(rows)
        source = client.get(reverse("userdata:download", args=[dataset.public_id, "source"]))
        assert source.status_code == 200
        assert "attachment" in source["Content-Disposition"]

    def test_document_has_own_source(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        record = _version(dataset).series.order_by("order").first()
        assert record is not None
        response = client.get(
            reverse("exports:document"),
            {"kind": "series", "series": _key(dataset, record), "format": "xlsx"},
        )
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(response.content))
        meta = [str(cell.value) for row in book["Отчёт"].iter_rows() for cell in row if cell.value]
        assert any("загружены пользователем" in value for value in meta)
        assert not any("Если быть точным" in value and "CC BY" in value for value in meta)
        pdf = client.get(
            reverse("exports:document"),
            {"kind": "ranking", "series": _key(dataset, record), "format": "pdf"},
        )
        assert pdf.status_code == 200
        assert pdf.content.startswith(b"%PDF")

    def test_delete_removes_files(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        directory = dataset.directory
        assert directory.exists()
        response = client.post(
            reverse("userdata:delete", args=[dataset.public_id]), {"confirm": "1"}
        )
        assert response.status_code == 302
        assert not Dataset.objects.filter(pk=dataset.pk).exists()
        assert not directory.exists()

    def test_delete_needs_confirmation(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        client.post(reverse("userdata:delete", args=[dataset.public_id]))
        assert Dataset.objects.filter(pk=dataset.pk).exists()

    def test_example(self, client: Client, warehouse: Any) -> None:
        response = client.post(reverse("userdata:example"))
        assert response.status_code == 302
        assert "/own-data/studies/" in response["Location"]
        dataset = Dataset.objects.get()
        assert dataset.source_url == "https://tochno.st/datasets/environment"
        assert _version(dataset).regions_count == 85
        assert client.get(response["Location"]).status_code == 200


@pytest.mark.django_db
class TestAgencyExamples:
    """Примеры ведомств (data/examples): открываются в мастере и скачиваются как опубликованы."""

    def test_every_example_is_listed_and_present(self, client: Client) -> None:
        from apps.userdata.examples import examples

        items = examples()
        assert len(items) >= 5
        body = client.get(reverse("userdata:index")).content.decode()
        for item in items:
            assert item.path.is_file(), item.file
            assert reverse("userdata:example-file", args=[item.slug]) in body

    def test_example_opens_in_the_wizard(self, client: Client, warehouse: Any) -> None:
        from apps.userdata.examples import example

        item = example("sfr-pension-supplement-2023")
        assert item is not None
        response = client.post(reverse("userdata:example-file", args=[item.slug]))
        dataset = Dataset.objects.get()
        assert response["Location"] == reverse("userdata:file", args=[dataset.public_id])
        assert dataset.title == f"Пример: {item.title}"
        assert dataset.source_url == item.source_url

    def test_unfinished_example_is_continued(self, client: Client, warehouse: Any) -> None:
        url = reverse("userdata:example-file", args=["sfr-pension-supplement-2023"])
        first = client.post(url)
        # Повторное нажатие на не собранный пример ведёт в тот же черновик, а не в копию.
        second = client.post(url)
        dataset = Dataset.objects.get()
        assert second["Location"] == first["Location"]
        listing = client.get(reverse("userdata:index")).content.decode()
        assert "не собрана" in listing
        assert reverse("userdata:dataset", args=[dataset.public_id]) in listing

    def test_example_file_downloads(self, client: Client) -> None:
        response = client.get(reverse("userdata:example-file", args=["fts-export-2021"]))
        assert response.status_code == 200
        assert "attachment" in response["Content-Disposition"]
        assert not Dataset.objects.exists()

    def test_unknown_example_is_not_found(self, client: Client) -> None:
        url = reverse("userdata:example-file", args=["nothing-here"])
        assert client.get(url).status_code == 404
        assert client.post(url).status_code == 404


@pytest.mark.django_db
class TestGuestsAndLimits:
    """Гость: три таблицы на сутки, перенос при входе; пределы частоты и места."""

    def test_guest_limit(self, client: Client, settings: Any, warehouse: Any) -> None:
        settings.USERDATA_GUEST_MAX_DATASETS = 1
        _upload(client, "environment.csv")
        response = client.post(
            reverse("userdata:upload"),
            {"action": "upload", "file": SimpleUploadedFile("a.csv", b"x;y\n1;2\n")},
        )
        assert "Без входа можно держать не больше 1 таблиц" in response.text
        assert Dataset.objects.count() == 1

    def test_rate_limit(self, client: Client, settings: Any, warehouse: Any) -> None:
        settings.USERDATA_GUEST_UPLOADS_PER_HOUR = 1
        _upload(client, "environment.csv")
        response = client.post(
            reverse("userdata:upload"),
            {"action": "upload", "file": SimpleUploadedFile("a.csv", b"x;y\n1;2\n")},
        )
        assert "Слишком много загрузок подряд" in response.text

    def test_space_quota(self, member_client: Client, settings: Any, warehouse: Any) -> None:
        _upload(member_client, "environment.csv")
        settings.USERDATA_QUOTA_BYTES = 10
        response = member_client.post(
            reverse("userdata:upload"),
            {"action": "upload", "file": SimpleUploadedFile("a.csv", b"x;y\n1;2\n")},
        )
        assert "Место для таблиц закончилось" in response.text

    def test_guest_tables_move_to_account(
        self, client: Client, make_user: Any, warehouse: Any
    ) -> None:
        dataset = _built(client)
        user = make_user(email="guest@example.com")
        client.force_login(user)
        dataset.refresh_from_db()
        assert dataset.owner == user
        assert dataset.expires_at is None
        assert client.get(reverse("userdata:dataset", args=[dataset.public_id])).status_code == 200

    def test_prune(self, client: Client, warehouse: Any, userdata_dir: Path) -> None:  # noqa: F811
        from django.core.management import call_command

        dataset = _built(client)
        Dataset.objects.filter(pk=dataset.pk).update(expires_at=timezone.now() - timedelta(hours=1))
        orphan = userdata_dir / "00000000-0000-0000-0000-000000000000"
        orphan.mkdir(parents=True)
        call_command("prune_personal_data", stdout=io.StringIO())
        assert not Dataset.objects.filter(pk=dataset.pk).exists()
        assert not orphan.exists()
        assert not dataset.directory.exists()


@pytest.mark.django_db
class TestAccount:
    """Таблицы в кабинете: «Обзор», выгрузка «Персональных данных», удаление учётной записи."""

    def test_overview_and_export(self, member_client: Client, member: Any, warehouse: Any) -> None:
        dataset = _built(member_client)
        overview = member_client.get(reverse("accounts:dashboard"))
        assert dataset.title in overview.text
        export = member_client.get(reverse("accounts:data-export"))
        payload = export.json()
        assert payload["own_data"][0]["title"] == dataset.title
        assert payload["own_data"][0]["series"]
        data = member_client.get(reverse("accounts:settings"))
        assert "Персональные данные" in data.text

    def test_account_deletion_removes_tables(
        self, member_client: Client, member: Any, warehouse: Any
    ) -> None:
        from apps.accounts.services import delete_account

        dataset = _built(member_client)
        directory = dataset.directory
        delete_account(member)
        assert not Dataset.objects.filter(pk=dataset.pk).exists()
        assert not directory.exists()


@pytest.mark.django_db
class TestMasks:
    """Коды-маски «Если быть точным» (8888 — «в регионе не ведутся наблюдения») — пропуски."""

    def test_codes_become_missing(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        version = _version(dataset)
        masks = version.report["extract"]["masks"]
        assert {"value": "8888", "count": 73, "detected": True, "masked": True} in masks
        key = _key(dataset, version.series.order_by("order").first())
        with scope.reading_as(fingerprint=dataset.guest_key):
            values = queries.series_values([key], ["RU-NEN", "RU-AD", "RU-KL"])
        assert all(
            value is None or value <= 100
            for years in values.get(key, {}).values()
            for value in years.values()
        )
        page = client.get(reverse("userdata:dataset", args=[dataset.public_id]), {"tab": "checks"})
        assert "Коды вместо чисел считаются пропусками" in page.text

    def test_person_can_keep_codes(self, client: Client, warehouse: Any) -> None:
        dataset = _built(client)
        url = reverse("userdata:series", args=[dataset.public_id])
        response = client.post(url, {"action": "masks"})
        assert response.status_code == 302
        assert _version(dataset).recipe["masks"] == []
        response = _build(client, dataset)
        assert response.status_code == 302
        version = _version(dataset)
        assert not any(item["masked"] for item in version.report["extract"]["masks"])
        key = _key(dataset, version.series.order_by("order").first())
        with scope.reading_as(fingerprint=dataset.guest_key):
            values = queries.series_values([key], ["RU-NEN", "RU-AD", "RU-KL"])
        assert any(
            value == 8888 for years in values.get(key, {}).values() for value in years.values()
        )


@pytest.mark.django_db
class TestIndicatorCodes:
    """Одно название у нескольких кодов показателя — разные показатели, а не повторы."""

    def test_same_name_different_codes(self, client: Client, warehouse: Any) -> None:
        lines = ["indicator_name;indicator_code;object_name;year;indicator_value"]
        names = ("Москва", "Республика Татарстан", "Свердловская область", "Омская область")
        for number, region in enumerate(names, start=1):
            for code, factor in (("Y331", 1), ("Y332", 10)):
                lines.append(f"Пробы почв с превышением, %;{code};{region};2021;{number * factor}")
        dataset = _upload(client, "soil.csv", "\n".join(lines).encode("utf-8"))
        _describe(client, dataset)
        assert _build(client, dataset, kind="relative", per=()).status_code == 302
        version = _version(dataset)
        assert version.report["extract"]["conflicts"] == 0
        titles = sorted(version.series.values_list("indicator", flat=True))
        assert titles == [
            "Пробы почв с превышением, % (Y331)",
            "Пробы почв с превышением, % (Y332)",
        ]
