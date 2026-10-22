"""
Отчёт исследования файлом: Word с карточками, заметками и картинками со страницы, Excel
с данными карточек; «не в отчёте»; доступ владельца и читателя закрытой ссылки.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from apps.userdata import studies
from apps.userdata.models import Study
from tests.support.documents import png
from tests.support.userdata import built as _built
from tests.support.userdata import first_record as _first
from tests.support.userdata import key_of as _key
from tests.support.userdata import userdata_dir  # noqa: F401 — приспособление модуля

pytestmark = [pytest.mark.integration, pytest.mark.django_db]

WORD = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
LINK = re.compile(r'value="(http://testserver/s/[A-Za-z0-9_-]+)"')


def _study(client: Client) -> Study:
    """Исследование с картой, рейтингом и заметкой."""
    dataset = _built(client)
    key = _key(dataset, _first(dataset))
    client.post(reverse("userdata:study-new"), {"title": "Воздух в городах"})
    study = Study.objects.latest("created_at")
    edit = reverse("userdata:study-edit", args=[study.public_id])
    client.post(edit, {"action": "add", "do": "view:map", "series": key})
    client.post(edit, {"action": "add", "do": "view:rankings", "series": key})
    client.post(edit, {"action": "add-text", "text": "Вывод: в столицах хуже."})
    study.refresh_from_db()
    return study


def _word(response: Any) -> str:
    """Текст тела документа Word из ответа."""
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
        return archive.read("word/document.xml").decode("utf-8")


def _url(study: Study) -> str:
    return reverse("userdata:study-report", args=[study.public_id])


class TestOwner:
    """Владелец скачивает Word и Excel; без сценария — Word без картинок."""

    def test_word_has_cards_notes_and_sources(self, member_client: Client, warehouse: Any) -> None:
        study = _study(member_client)
        response = member_client.get(_url(study), {"format": "docx"})
        assert response.status_code == 200
        assert response["Content-Type"] == WORD
        assert "filename*=UTF-8''" in response["Content-Disposition"]
        assert response["Cache-Control"] == "private, no-store"
        body = _word(response)
        assert "Воздух в городах" in body
        assert "Вывод: в столицах хуже." in body
        assert "Население в городах с высоким" in body
        # Рейтинг — таблицей, источники — списком в конце.
        assert "<w:tbl>" in body and "Источники" in body
        assert "<w:drawing>" not in body

    def test_pictures_from_the_page_are_embedded(
        self, member_client: Client, warehouse: Any
    ) -> None:
        """Картинка карты — PNG из формы; не PNG и чужие поля отбрасываются молча."""
        study = _study(member_client)
        block = studies.blocks_of(study)[0]
        response = member_client.post(
            _url(study),
            {
                "format": "docx",
                f"image-{block['id']}-0": SimpleUploadedFile("map.png", png(40, 20)),
                f"image-{block['id']}-1": SimpleUploadedFile("map.gif", b"GIF89a" + bytes(40)),
                "image-nothing-0": SimpleUploadedFile("x.png", png(4, 4)),
            },
        )
        assert response.status_code == 200
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            media = [name for name in archive.namelist() if name.startswith("word/media/")]
        assert media == ["word/media/image1.png"]

    def test_excel_has_a_sheet_per_card_with_data(
        self, member_client: Client, warehouse: Any
    ) -> None:
        study = _study(member_client)
        response = member_client.get(_url(study), {"format": "xlsx"})
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(response.content))
        assert book.sheetnames == ["Отчёт", "1. Карта", "2. Рейтинг"]
        sheet = book["1. Карта"]
        assert sheet["A1"].value.startswith("Население в городах")
        header = [cell.value for cell in sheet[3]]
        assert header[:4] == ["Субъект", "Код", "Показатель", "Единица"]
        assert "2014" in header

    def test_unknown_format_is_not_found(self, member_client: Client, warehouse: Any) -> None:
        study = _study(member_client)
        assert member_client.get(_url(study), {"format": "pdf"}).status_code == 404

    def test_card_left_out_of_the_report(self, member_client: Client, warehouse: Any) -> None:
        """«Не включать в отчёт» — пометка на поле, нет в Word, Excel и печати."""
        study = _study(member_client)
        note, ranking = studies.blocks_of(study)[2], studies.blocks_of(study)[1]
        edit = reverse("userdata:study-edit", args=[study.public_id])
        for block in (note, ranking):
            member_client.post(edit, {"action": "report", "block": block["id"]})
        study.refresh_from_db()
        assert [block.get("report", True) for block in studies.blocks_of(study)] == [
            True,
            False,
            False,
        ]
        page = member_client.get(reverse("userdata:study", args=[study.public_id]))
        assert page.text.count("is-off-report") == 2
        assert "Включить в отчёт" in page.text
        body = _word(member_client.get(_url(study), {"format": "docx"}))
        assert "Вывод: в столицах хуже." not in body and "<w:tbl>" not in body
        book = load_workbook(io.BytesIO(member_client.get(_url(study), {"format": "xlsx"}).content))
        assert book.sheetnames == ["Отчёт", "1. Карта"]
        # Повторная отметка возвращает карточку в отчёт.
        member_client.post(edit, {"action": "report", "block": note["id"]})
        study.refresh_from_db()
        assert "report" not in studies.blocks_of(study)[2]


class TestReader:
    """Читатель закрытой ссылки скачивает отчёт, только если ссылка это разрешает."""

    def _reader(self, member_client: Client, study: Study, *, downloads: bool) -> Client:
        data = {"days": "30", **({"downloads": "on"} if downloads else {})}
        member_client.post(reverse("userdata:study-share-create", args=[study.public_id]), data)
        page = member_client.get(reverse("userdata:study", args=[study.public_id]))
        found = LINK.findall(page.text)
        assert found
        reader = Client()
        reader.get(found[-1].removeprefix("http://testserver"))
        return reader

    def test_without_downloads(self, member_client: Client, warehouse: Any) -> None:
        study = _study(member_client)
        reader = self._reader(member_client, study, downloads=False)
        page = reader.get(reverse("userdata:study", args=[study.public_id]))
        assert page.status_code == 200 and "study-report" not in page.text
        assert reader.get(_url(study), {"format": "docx"}).status_code == 404

    def test_with_downloads(self, member_client: Client, warehouse: Any) -> None:
        study = _study(member_client)
        reader = self._reader(member_client, study, downloads=True)
        page = reader.get(reverse("userdata:study", args=[study.public_id]))
        assert 'id="study-report"' in page.text
        assert reader.get(_url(study), {"format": "xlsx"}).status_code == 200
        # Форма Word со страницы читателя — с ключом CSRF, как у владельца.
        assert "csrfmiddlewaretoken" in page.text.split('id="study-report"', 1)[1]

    def test_stranger(self, member_client: Client, warehouse: Any) -> None:
        study = _study(member_client)
        assert Client().get(_url(study), {"format": "docx"}).status_code == 404
