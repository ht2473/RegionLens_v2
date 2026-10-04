"""Вывод отчёта в Excel: числа — числами с форматом ячейки, реквизиты — на первом листе."""

from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..reports.base import COLUMN_INTEGER, COLUMN_NUMBER, COLUMN_PERCENT, ReportDocument, Table
from .palette import document_colours
from .safety import safe_cell

# Форматы чисел по видам столбцов.
NUMBER_FORMATS = {
    COLUMN_NUMBER: "# ##0.00",
    COLUMN_INTEGER: "# ##0",
    COLUMN_PERCENT: "0.0 %",
}

# Ограничение Excel на длину имени листа.
SHEET_NAME_LIMIT = 31

# Ширина первого столбца листа реквизитов.
META_LABEL_WIDTH = 28
META_VALUE_WIDTH = 80


def render(document: ReportDocument) -> bytes:
    """Собрать книгу Excel по отчёту."""
    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_summary_sheet(workbook, document)

    used_names: set[str] = {"Отчёт"}
    for index, table in enumerate(document.tables, start=1):
        name = _unique_sheet_name(table.title or f"Таблица {index}", used_names)
        used_names.add(name)
        _write_table_sheet(workbook.create_sheet(name), table)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_summary_sheet(workbook: Workbook, document: ReportDocument) -> None:
    """Записать лист с названием отчёта и его реквизитами."""
    palette = document_colours()
    sheet = workbook.create_sheet("Отчёт")
    sheet["A1"] = safe_cell(document.title)
    sheet["A1"].font = Font(bold=True, size=14, color=palette.text)
    sheet["A2"] = safe_cell(document.subtitle)
    sheet["A2"].font = Font(color=palette.muted)

    row = 4
    for label, value in document.meta:
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True, color=palette.text)
        sheet.cell(row=row, column=2, value=safe_cell(value))
        row += 1

    if document.footer:
        sheet.cell(row=row + 1, column=1, value=document.footer)

    sheet.column_dimensions["A"].width = META_LABEL_WIDTH
    sheet.column_dimensions["B"].width = META_VALUE_WIDTH


def _write_table_sheet(sheet: Worksheet, table: Table) -> None:
    """Записать одну таблицу отчёта на лист; шапка — как у таблиц сайта."""
    palette = document_colours()
    rule = Side(style="thin", color=palette.rule)
    border = Border(left=rule, right=rule, top=rule, bottom=rule)
    header_border = Border(
        left=rule, right=rule, top=rule, bottom=Side(style="medium", color=palette.header_rule)
    )
    header_fill = PatternFill("solid", fgColor=palette.header)
    header_font = Font(bold=True, size=11, color=palette.header_text)

    header_row = 1
    if table.note:
        sheet.cell(row=1, column=1, value=safe_cell(table.note))
        header_row = 3

    for index, column in enumerate(table.columns, start=1):
        cell = sheet.cell(row=header_row, column=index, value=safe_cell(column.title))
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
        sheet.column_dimensions[get_column_letter(index)].width = column.width

    for offset, row in enumerate(table.rows, start=header_row + 1):
        for index, column in enumerate(table.columns, start=1):
            cell = sheet.cell(row=offset, column=index, value=safe_cell(row.get(column.key)))
            cell.border = border
            if column.is_numeric:
                cell.number_format = NUMBER_FORMATS.get(column.kind, "General")
                cell.alignment = Alignment(horizontal="right")

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=2)
    sheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(table.columns))}{header_row + len(table.rows)}"
    )


def _unique_sheet_name(title: str, used: set[str]) -> str:
    """Привести название таблицы к допустимому и неповторяющемуся имени листа (до 31 знака)."""
    cleaned = re.sub(r"[\\/*?:\[\]]", " ", title).strip() or "Таблица"
    name = cleaned[:SHEET_NAME_LIMIT]

    suffix = 2
    while name in used:
        tail = f" ({suffix})"
        name = cleaned[: SHEET_NAME_LIMIT - len(tail)] + tail
        suffix += 1
    return name
